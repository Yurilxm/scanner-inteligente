"""Atualiza Planilha_Contas.xlsx a partir dos CSVs exportados pelo app Control.

- Le todos os .csv da pasta, do mais antigo para o mais recente (o mais novo prevalece)
- Contas novas sao adicionadas; contas ja existentes sao atualizadas (pago, data de pagamento...)
- A chave de cada conta e (Perfil, ID), para nao misturar os perfis
- Faz backup da planilha antes de alterar
"""
import csv
import os
import shutil
import sys
from datetime import datetime

from openpyxl import Workbook, load_workbook

PASTA_ATUAL = os.path.dirname(os.path.abspath(__file__))
NOME_PLANILHA = "Planilha_Contas.xlsx"

COLUNAS = [
    "Perfil", "ID", "Tipo", "Emissor", "Referencia", "Vencimento",
    "Valor (R$)", "CNPJ", "Pago", "Data Pagamento", "Linha Digitavel",
]
COLUNAS_DATA = {"Vencimento", "Data Pagamento"}


def converter(coluna, texto):
    """Valor vira numero e datas viram datas, para o Excel conseguir somar, ordenar e filtrar."""
    texto = (texto or "").strip()
    if not texto:
        return None
    if coluna == "Valor (R$)":
        try:
            return float(texto.replace(".", "").replace(",", "."))
        except ValueError:
            return texto
    if coluna in COLUNAS_DATA:
        try:
            return datetime.strptime(texto, "%d/%m/%Y")
        except ValueError:
            return texto
    return texto


def ler_csvs(pasta):
    arquivos = [os.path.join(pasta, f) for f in os.listdir(pasta) if f.lower().endswith(".csv")]
    arquivos.sort(key=os.path.getmtime)
    linhas = []
    for caminho in arquivos:
        with open(caminho, "r", encoding="utf-8-sig", newline="") as f:
            for linha in csv.DictReader(f, delimiter=";"):
                if (linha.get("ID") or "").strip():
                    linhas.append(linha)
    return arquivos, linhas


def preparar_cabecalho(ws):
    """Mapeia coluna -> indice; acrescenta no fim as colunas que planilhas antigas nao tinham."""
    cabecalho = [c.value for c in ws[1]] if ws.max_row >= 1 else []
    if not any(cabecalho):
        cabecalho = []
    indices = {nome: i + 1 for i, nome in enumerate(cabecalho) if nome}
    for nome in COLUNAS:
        if nome not in indices:
            indices[nome] = max(indices.values(), default=0) + 1
            ws.cell(row=1, column=indices[nome], value=nome)
    return indices


def atualizar(pasta=PASTA_ATUAL):
    arquivos, linhas = ler_csvs(pasta)
    if not arquivos:
        print("\n[!] Nenhum arquivo CSV encontrado nesta pasta.")
        print("-> Exporte o CSV pelo app (Biblioteca > CSV > Compartilhar) e coloque nesta pasta.")
        return None

    print(f"\n[1/3] Lendo {len(arquivos)} arquivo(s): " + ", ".join(os.path.basename(a) for a in arquivos))

    planilha = os.path.join(pasta, NOME_PLANILHA)
    if os.path.exists(planilha):
        backup = planilha.replace(".xlsx", f"_backup_{datetime.now():%Y%m%d%H%M%S}.xlsx")
        shutil.copy2(planilha, backup)
        print(f"[2/3] Backup criado: {os.path.basename(backup)}")
        wb = load_workbook(planilha)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Contas"
        print("[2/3] Nova planilha criada.")

    col = preparar_cabecalho(ws)

    # Contas que ja estao na planilha, pela chave (Perfil, ID)
    existentes = {}
    for numero in range(2, ws.max_row + 1):
        conta_id = ws.cell(row=numero, column=col["ID"]).value
        if conta_id is None:
            continue
        perfil = ws.cell(row=numero, column=col["Perfil"]).value or ""
        existentes[(str(perfil), str(conta_id))] = numero

    # Conta cada conta uma vez, mesmo que apareca em mais de um CSV
    novas, atualizadas = set(), set()
    for linha in linhas:
        chave = ((linha.get("Perfil") or "").strip(), linha["ID"].strip())
        numero = existentes.get(chave)
        if numero is None and chave[0]:
            # Linha de versao antiga, gravada sem perfil: reaproveita em vez de duplicar
            numero = existentes.pop(("", chave[1]), None)
            if numero is not None:
                existentes[chave] = numero
        if numero is None:
            numero = ws.max_row + 1
            existentes[chave] = numero
            novas.add(chave)
        elif chave not in novas:
            atualizadas.add(chave)
        for nome in COLUNAS:
            valor = chave[0] if nome == "Perfil" else converter(nome, linha.get(nome))
            celula = ws.cell(row=numero, column=col[nome], value=valor)
            if nome == "Valor (R$)" and isinstance(valor, float):
                celula.number_format = "#,##0.00"
            elif nome in COLUNAS_DATA and isinstance(valor, datetime):
                celula.number_format = "DD/MM/YYYY"

    wb.save(planilha)
    print(f"[3/3] Pronto! {len(novas)} conta(s) nova(s), {len(atualizadas)} atualizada(s).")
    print(f"-> Arquivo salvo: {planilha}")
    return {"adicionadas": len(novas), "atualizadas": len(atualizadas)}


def main():
    print("=" * 40)
    print("   CONTROL - ATUALIZADOR DE PLANILHA")
    print("=" * 40)
    try:
        atualizar()
    except PermissionError:
        print("\n[!] Nao consegui salvar: feche a planilha no Excel e tente de novo.")
    input("\nAperte Enter para fechar...")


if __name__ == "__main__":
    sys.exit(main())
