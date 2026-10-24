"""Testes do atualizador da planilha. Rodar com: python -m unittest Controle_Excel/test_atualizar_planilha.py"""
import os
import sys
import tempfile
import time
import unittest
from datetime import datetime

from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import atualizar_planilha as ap  # noqa: E402

HEADER = "ID;Tipo;Emissor;Referencia;Vencimento;Valor (R$);CNPJ;Pago;Data Pagamento;Linha Digitavel;Perfil"


def escrever_csv(pasta, nome, linhas, header=HEADER):
    caminho = os.path.join(pasta, nome)
    with open(caminho, "w", encoding="utf-8-sig", newline="") as f:
        f.write(header + "\n" + "\n".join(linhas) + "\n")
    return caminho


def ler_planilha(pasta):
    ws = load_workbook(os.path.join(pasta, ap.NOME_PLANILHA)).active
    cab = [c.value for c in ws[1]]
    return [dict(zip(cab, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]


class AtualizarPlanilhaTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.pasta = self.tmp.name

    def tearDown(self):
        self.tmp.cleanup()

    def test_sem_csv_nao_cria_planilha(self):
        self.assertIsNone(ap.atualizar(self.pasta))
        self.assertFalse(os.path.exists(os.path.join(self.pasta, ap.NOME_PLANILHA)))

    def test_cria_planilha_com_valor_numerico_e_datas(self):
        escrever_csv(self.pasta, "relatorio_yuri.csv", ["1;BOLETO;ENEL;10/2026;10/11/2026;1234,56;;Nao;;;yuri"])
        self.assertEqual(ap.atualizar(self.pasta), {"adicionadas": 1, "atualizadas": 0})
        [linha] = ler_planilha(self.pasta)
        self.assertEqual(linha["Perfil"], "yuri")
        self.assertEqual(linha["Valor (R$)"], 1234.56)
        self.assertEqual(linha["Vencimento"], datetime(2026, 11, 10))
        self.assertIsNone(linha["Data Pagamento"])

    def test_conta_paga_depois_e_atualizada_sem_duplicar(self):
        escrever_csv(self.pasta, "relatorio_yuri.csv", ["1;BOLETO;ENEL;10/2026;10/11/2026;181,31;;Nao;;;yuri"])
        ap.atualizar(self.pasta)
        time.sleep(0.05)
        escrever_csv(self.pasta, "relatorio_yuri.csv", ["1;BOLETO;ENEL;10/2026;10/11/2026;181,31;;Sim;08/11/2026;;yuri"])
        self.assertEqual(ap.atualizar(self.pasta), {"adicionadas": 0, "atualizadas": 1})
        [linha] = ler_planilha(self.pasta)
        self.assertEqual(linha["Pago"], "Sim")
        self.assertEqual(linha["Data Pagamento"], datetime(2026, 11, 8))

    def test_perfis_com_mesmo_id_nao_se_misturam(self):
        escrever_csv(self.pasta, "relatorio_yuri.csv", ["1;BOLETO;ENEL;;;100,00;;Nao;;;yuri"])
        escrever_csv(self.pasta, "relatorio_mae.csv", ["1;BOLETO;PROLAGOS;;;200,00;;Nao;;;mae"])
        ap.atualizar(self.pasta)
        linhas = {l["Perfil"]: l["Emissor"] for l in ler_planilha(self.pasta)}
        self.assertEqual(linhas, {"yuri": "ENEL", "mae": "PROLAGOS"})

    def test_csv_mais_recente_prevalece(self):
        antigo = escrever_csv(self.pasta, "antigo.csv", ["1;BOLETO;ENEL;;;100,00;;Sim;01/11/2026;;yuri"])
        novo = escrever_csv(self.pasta, "novo.csv", ["1;BOLETO;ENEL;;;100,00;;Nao;;;yuri"])
        os.utime(antigo, (1_000_000, 1_000_000))
        os.utime(novo, (2_000_000, 2_000_000))
        self.assertEqual(ap.atualizar(self.pasta), {"adicionadas": 1, "atualizadas": 0})
        [linha] = ler_planilha(self.pasta)
        self.assertEqual(linha["Pago"], "Nao")

    def test_planilha_antiga_ganha_colunas_e_linhas_sem_perfil_sao_reaproveitadas(self):
        # Formato da versao anterior: sem Perfil e sem Data Pagamento, valor como texto
        wb = Workbook()
        ws = wb.active
        ws.append(["ID", "Tipo", "Emissor", "Referencia", "Vencimento", "Valor (R$)", "CNPJ", "Pago", "Linha Digitavel"])
        ws.append(["doc_001", "BOLETO", "ENEL", "10/2026", "10/11/2026", "181,31", "", "Nao", ""])
        ws.append(["anotacao manual", "OUTROS", "", "", "", "", "", "", ""])
        wb.save(os.path.join(self.pasta, ap.NOME_PLANILHA))

        escrever_csv(self.pasta, "relatorio_yuri.csv", ["doc_001;BOLETO;ENEL;10/2026;10/11/2026;181,31;;Sim;09/11/2026;;yuri"])
        self.assertEqual(ap.atualizar(self.pasta), {"adicionadas": 0, "atualizadas": 1})

        linhas = ler_planilha(self.pasta)
        self.assertEqual(len(linhas), 2)
        self.assertEqual(linhas[0]["Perfil"], "yuri")
        self.assertEqual(linhas[0]["Pago"], "Sim")
        self.assertEqual(linhas[0]["Data Pagamento"], datetime(2026, 11, 9))
        self.assertEqual(linhas[1]["ID"], "anotacao manual")  # linha que nao veio do app fica intacta
        backups = [f for f in os.listdir(self.pasta) if "_backup_" in f]
        self.assertEqual(len(backups), 1)

    def test_csv_antigo_sem_coluna_perfil(self):
        escrever_csv(
            self.pasta, "relatorio_documentos.csv",
            ["doc_001;BOLETO;ENEL;10/2026;10/11/2026;181,31;61.695.227/0001-93;Nao;"],
            header="ID;Tipo;Emissor;Referencia;Vencimento;Valor (R$);CNPJ;Pago;Linha Digitavel",
        )
        self.assertEqual(ap.atualizar(self.pasta), {"adicionadas": 1, "atualizadas": 0})
        [linha] = ler_planilha(self.pasta)
        self.assertIsNone(linha["Perfil"])
        self.assertEqual(linha["CNPJ"], "61.695.227/0001-93")

    def test_campos_com_ponto_e_virgula_entre_aspas(self):
        escrever_csv(self.pasta, "r.csv", ['1;BOLETO;"ENEL; DISTRIBUICAO";;;10,00;;Nao;;;yuri'])
        ap.atualizar(self.pasta)
        [linha] = ler_planilha(self.pasta)
        self.assertEqual(linha["Emissor"], "ENEL; DISTRIBUICAO")


if __name__ == "__main__":
    unittest.main()
